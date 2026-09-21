from openpyxl import load_workbook
from models import Student
from auth import hash_password
from sqlalchemy.orm import Session
import json

REQUIRED_FIELDS = ("register_no", "name", "dob", "department", "year", "section")
HEADER_ALIASES = {
    "register no": "register_no",
    "register number": "register_no",
    "reg no": "register_no",
    "reg number": "register_no",
    "date of birth": "dob",
    "dept": "department",
    "class year": "year",
}


def _canonical_field(field_name: str):
    if field_name is None:
        return None
    normalized = str(field_name).strip().lower().replace("-", " ").replace("_", " ")
    alias = HEADER_ALIASES.get(normalized)
    if alias:
        return alias
    normalized = normalized.replace(" ", "_")
    if normalized in REQUIRED_FIELDS:
        return normalized
    return None


def _to_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _normalize_student_payload(raw_data):
    normalized = {}
    missing = []

    for field in REQUIRED_FIELDS:
        value = _to_text(raw_data.get(field))
        if not value:
            missing.append(field)
        normalized[field] = value

    return normalized, missing


def _import_student_records(records, db: Session):
    summary = {
        "status": "success",
        "total_rows": len(records),
        "imported": 0,
        "skipped_duplicates": 0,
        "skipped_invalid": 0,
        "invalid_rows": [],
    }

    try:
        for index, record in enumerate(records, start=1):
            student_data, missing = _normalize_student_payload(record)
            if missing:
                summary["skipped_invalid"] += 1
                summary["invalid_rows"].append({"row": index, "reason": f"Missing required fields: {', '.join(missing)}"})
                continue

            if db.query(Student).filter(Student.register_no == student_data["register_no"]).first():
                summary["skipped_duplicates"] += 1
                continue

            student = Student(
                register_no=student_data["register_no"],
                name=student_data["name"],
                dob=student_data["dob"],
                department=student_data["department"],
                year=student_data["year"],
                section=student_data["section"],
                password_hash=hash_password(student_data["dob"]),
            )
            db.add(student)
            summary["imported"] += 1

        db.commit()
        if len(summary["invalid_rows"]) > 20:
            summary["invalid_rows"] = summary["invalid_rows"][:20]

        summary["message"] = (
            f"Imported {summary['imported']} of {summary['total_rows']} rows "
            f"({summary['skipped_duplicates']} duplicates, {summary['skipped_invalid']} invalid)."
        )
        return summary
    except Exception as e:
        db.rollback()
        return {"status": "error", "message": str(e)}


def import_students_from_excel(file_path: str, db: Session):
    try:
        workbook = load_workbook(file_path)
        sheet = workbook.active

        raw_headers = [cell.value for cell in sheet[1]]
        mapped_headers = [_canonical_field(header) for header in raw_headers]

        missing_headers = [field for field in REQUIRED_FIELDS if field not in mapped_headers]
        if missing_headers:
            raise ValueError(f"Excel missing required columns: {', '.join(missing_headers)}")

        records = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_data = {}
            for idx, value in enumerate(row):
                if idx >= len(mapped_headers):
                    continue
                field = mapped_headers[idx]
                if field:
                    row_data[field] = value
            if any(_to_text(value) for value in row_data.values()):
                records.append(row_data)

        return _import_student_records(records, db)
    except Exception as e:
        db.rollback()
        return {"status": "error", "message": str(e)}


def import_students_from_json(file_path: str, db: Session):
    try:
        with open(file_path, "r", encoding="utf-8") as file:
            students_data = json.load(file)

        if not isinstance(students_data, list):
            raise ValueError("JSON file must contain an array of student objects")

        records = []
        for item in students_data:
            if not isinstance(item, dict):
                records.append({})
                continue

            normalized = {}
            for key, value in item.items():
                field = _canonical_field(key)
                if field:
                    normalized[field] = value
            records.append(normalized)

        return _import_student_records(records, db)
    except Exception as e:
        db.rollback()
        return {"status": "error", "message": str(e)}
